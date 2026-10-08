"""Контрольна перевірка пакета (на перерахованому файлі з кешованими значеннями)."""
import sys, re, os, json
import openpyxl
sys.path.insert(0, os.path.dirname(__file__))
from num2words_uk import words, fmt_money

path = sys.argv[1]
items = json.load(open(sys.argv[2], encoding="utf-8")) if len(sys.argv) > 2 else None
wf = openpyxl.load_workbook(path)
wv = openpyxl.load_workbook(path, data_only=True)
errs = []
for ws in wv:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("#"):
                errs.append(f"{ws.title}!{c.coordinate} = {c.value}")
for ws in wf:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and ("[" in c.value and c.value.startswith("=") or "#REF" in c.value):
                errs.append(f"{ws.title}!{c.coordinate} підозріла формула {c.value}")

def block(ws, start, numcol):
    r, nums = start, []
    while isinstance(ws[f"{numcol}{r}"].value, (int, float)):
        nums.append(ws[f"{numcol}{r}"].value); r += 1
    return nums, r

cfg = {"рахунок": (20, "A", "G"), "накладна": (17, "A", "G"), "Дрюк": (13, "A", "J"), "Нагорнюк": (13, "A", "J"),
       "Лозко": (13, "A", "J"), "ТОВ Колотір": (13, "A", "J"), "пропозиція 1": (16, "A", "G"),
       "пропозиція 2": (6, "A", "H"), "пропозиція 3": (14, "A", "G")}
mult = {"рахунок": 1, "накладна": 1, "Дрюк": 1, "Нагорнюк": 1.05, "Лозко": 1.08, "ТОВ Колотір": 1.1,
        "пропозиція 1": 1.05, "пропозиція 2": 1.08, "пропозиція 3": 1.1}
rah = wv["рахунок"]
base = []
r = 20
while isinstance(rah[f"A{r}"].value, (int, float)):
    base.append((rah[f"E{r}"].value, rah[f"F{r}"].value)); r += 1
n = len(base)
if items:
    for k, it in enumerate(items):
        row = 20 + k
        got = (rah[f"B{row}"].value, str(rah[f"C{row}"].value), rah[f"E{row}"].value, rah[f"F{row}"].value)
        exp = (it["name"], it["article"], it["qty"], it["price"])
        if got != exp:
            errs.append(f"рахунок рядок {row}: {got} ≠ {exp}")
print(f"{'аркуш':14} {'рядків':>6} {'сума':>14} {'очікувано':>14}")
for name, (start, nc, sc) in cfg.items():
    ws = wv[name]
    nums, end = block(ws, start, nc)
    if nums != list(range(1, n + 1)):
        errs.append(f"{name}: нумерація/кількість рядків {len(nums)} ≠ {n}")
    s = sum(ws[f"{sc}{r}"].value or 0 for r in range(start, end))
    exp = sum(q * round(p * mult[name] + 1e-9) for q, p in base)
    print(f"{name:14} {len(nums):>6} {fmt_money(s):>14} {fmt_money(exp):>14}")
    if abs(s - exp) > 0.001:
        errs.append(f"{name}: сума {s} ≠ {exp}")
    # підсумкові клітинки
    for row in ws.iter_rows(min_row=end):
        for c in row:
            if isinstance(c.value, (int, float)) and c.value > 1000 and abs(c.value - exp) > 0.001 and c.column_letter in (sc, "F", "H"):
                errs.append(f"{name}!{c.coordinate} = {c.value}, очікувано {exp}")
    # тексти
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not isinstance(v, str):
                continue
            if v.startswith("Сума прописом"):
                ok = v in (f"Сума прописом {words(exp)}", f"Сума прописом {words(exp, False)}")
                if not ok: errs.append(f"{name}!{c.coordinate} прописом: {v}")
            if v.startswith("Загальна вартість") and v != f"Загальна вартість (ціна): {fmt_money(exp)} грн ({words(exp)}) без ПДВ.":
                errs.append(f"{name}!{c.coordinate}: {v}")
            if v.startswith("Всього на суму") and v != f"Всього на суму: {words(exp)}":
                errs.append(f"{name}!{c.coordinate}: {v}")
            if v.startswith("ПДВ:") and re.search(r"\d", v) and not re.search(r"0[.,]00", v):
                errs.append(f"{name}!{c.coordinate}: {v}")
    if name == "Дрюк":
        for r in range(start, end):
            u = ws[f"N{r}"].value
            if not re.fullmatch(r"\d{10}", str(u or "")):
                errs.append(f"Дрюк!N{r}: УКТЗЕД '{u}'")
            if not ws[f"L{r}"].value:
                print(f"  ! Дрюк!L{r}: виробник не вказано ({ws[f'B{r}'].value} {ws[f'F{r}'].value})")
k = wv["калькуляція"]
tot = sum(q * p for q, p in base)
print("калькуляція: закупівля", round(k["C16"].value, 2), "інші", round(k["C20"].value, 2),
      "прибуток", round(k["C19"].value, 2), "всього", k["C22"].value)
if abs(k["C22"].value - tot) > 0.01:
    errs.append(f"калькуляція всього {k['C22'].value} ≠ {tot}")
print("A1:", rah["A1"].value); print("D2:", wv["накладна"]["D2"].value.replace("\n", " | "))
print("\nПОМИЛКИ:" if errs else "\nПомилок не знайдено.", *errs, sep="\n")
