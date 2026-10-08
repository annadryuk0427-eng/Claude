"""Заповнення шаблону пакета документів на запчастини (openpyxl, без Excel).

python scripts/fill_invoice.py --src template.xlsx --items items.json --dst out.xlsx --date ДД.ММ.РРРР
"""
import argparse, copy, json, re, sys, os
import openpyxl
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(__file__))
from num2words_uk import words, fmt_money

TPL_N = 21  # позицій у шаблоні
# аркуш: (перший рядок позицій, генератор формул/значень для позиції k (0-based) -> {col: value})
R0 = 20  # перший рядок позицій на «рахунок»
D0 = 13  # на «Дрюк»


def rah(k, it):
    r = R0 + k
    return {"A": k + 1, "B": it["name"], "C": it["article"], "D": "шт", "E": it["qty"],
            "F": it["price"], "G": f"=F{r}*E{r}"}


def nakl(k, it):
    s = R0 + k
    return {c: f"=рахунок!{c}{s}" for c in "ABCDEFG"}


def druk(k, it):
    s = R0 + k
    return {"A": f"=рахунок!A{s}", "B": f"=рахунок!B{s}", "F": f"=рахунок!C{s}", "H": f"=рахунок!E{s}",
            "I": f"=рахунок!F{s}", "J": f"=рахунок!G{s}", "K": "Новий",
            "L": it.get("manufacturer") or None, "M": it.get("country") or None,
            "N": int(it["uktzed"]) if it.get("uktzed") else None}


def supplier(mult):
    def f(k, it, row):
        d = D0 + k
        return {"A": f"=Дрюк!A{d}", "B": f"=Дрюк!B{d}", "F": f"=Дрюк!F{d}", "H": f"=Дрюк!H{d}",
                "I": f"=ROUND(Дрюк!I{d}*{mult},0)", "J": f"=H{row}*I{row}", "K": "Новий",
                "L": f"=Дрюк!L{d}", "M": f"=Дрюк!M{d}", "N": f"=Дрюк!N{d}"}
    return f


def prop1(k, it, row):
    s = R0 + k
    return {"A": f"=рахунок!A{s}", "B": f"=рахунок!B{s}", "D": f"=рахунок!D{s}", "E": f"=рахунок!E{s}",
            "F": f"=ROUND((рахунок!F{s}*1.05),0)", "G": f"=F{row}*E{row}"}


def prop2(k, it, row):
    s = R0 + k
    return {"A": f"=рахунок!A{s}", "B": f"=рахунок!B{s}", "E": f"=рахунок!D{s}", "F": f"=рахунок!E{s}",
            "G": f"=ROUND((рахунок!F{s}*1.08),0)", "H": f"=G{row}*F{row}"}


def prop3(k, it, row):
    s = R0 + k
    return {"A": f"=рахунок!A{s}", "B": f"=рахунок!B{s}", "C": f"=рахунок!C{s}", "D": f"=рахунок!D{s}",
            "E": f"=рахунок!E{s}", "F": f"=ROUND((рахунок!F{s}*1.1),0)", "G": f"=F{row}*E{row}"}


SHEETS = {
    "рахунок": (20, lambda k, it, row: rah(k, it)),
    "накладна": (17, lambda k, it, row: nakl(k, it)),
    "Дрюк": (13, lambda k, it, row: druk(k, it)),
    "Нагорнюк": (13, supplier(1.05)),
    "Лозко": (13, supplier(1.08)),
    "ТОВ Колотір": (13, supplier(1.1)),
    "пропозиція 1": (16, prop1),
    "пропозиція 2": (6, prop2),
    "пропозиція 3": (14, prop3),
}
LAST = {name: start + TPL_N - 1 for name, (start, _) in SHEETS.items()}
REF = re.compile(r"(?:(рахунок|накладна|Дрюк|Нагорнюк|Лозко|калькуляція|'[^']+')!)?(\$?)([A-Z]{1,2})(\$?)(\d+)(?![\d(])")


def shift_formula(f, own, delta):
    def rep(m):
        sh, d1, col, d2, row = m.groups()
        target = (sh or own).strip("'")
        last = LAST.get(target)
        row = int(row)
        if last is not None and (row > last or (row == last and f[m.start() - 1:m.start()] == ":")):
            row += delta
        return f"{sh + '!' if sh else ''}{d1}{col}{d2}{row}"
    return REF.sub(rep, f)


def resize(ws, start, n, gen, items):
    last = start + TPL_N - 1
    delta = n - TPL_N
    maxc = ws.max_column
    tail_merges = [r for r in ws.merged_cells.ranges if r.min_row > last]
    item_merges = [r for r in ws.merged_cells.ranges if start <= r.min_row <= last]
    pattern = [(r.min_col, r.max_col) for r in item_merges if r.min_row == start]
    for r in tail_merges + item_merges:
        ws.unmerge_cells(str(r))
    heights = {r: ws.row_dimensions[r].height for r in range(last + 1, ws.max_row + 1)}
    style_row = {c: copy.copy(ws.cell(start, c)._style) for c in range(1, maxc + 1)}
    h_item = ws.row_dimensions[start].height
    if delta > 0:
        ws.move_range(f"A{last + 1}:{get_column_letter(maxc)}{ws.max_row}", rows=delta)
    elif delta < 0:
        for r in range(start + n, last + 1):
            for c in range(1, maxc + 1):
                ws.cell(r, c).value = None
        ws.move_range(f"A{last + 1}:{get_column_letter(maxc)}{ws.max_row}", rows=delta)
    for r in list(heights):
        ws.row_dimensions[r].height = None
    for r, h in heights.items():
        ws.row_dimensions[r + delta].height = h
    for k, it in enumerate(items):
        row = start + k
        ws.row_dimensions[row].height = h_item
        for c in range(1, maxc + 1):
            ws.cell(row, c)._style = copy.copy(style_row[c])
            ws.cell(row, c).value = None
        for col, v in gen(k, it, row).items():
            ws[f"{col}{row}"] = v
        for c1, c2 in pattern:
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    for r in tail_merges:
        ws.merge_cells(start_row=r.min_row + delta, start_column=r.min_col,
                       end_row=r.max_row + delta, end_column=r.max_col)


def find(ws, prefix):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith(prefix):
                return c
    raise KeyError(f"{ws.title}: '{prefix}' не знайдено")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--items", required=True)
    ap.add_argument("--dst", required=True)
    ap.add_argument("--date", required=True)
    a = ap.parse_args()
    if os.path.exists(a.dst):
        sys.exit(f"{a.dst} уже існує — не перезаписую")
    items = json.load(open(a.items, encoding="utf-8"))
    n, delta = len(items), len(items) - TPL_N
    wb = openpyxl.load_workbook(a.src)
    # 1) зсув посилань у формулах поза блоками позицій (підсумки, калькуляція)
    for ws in wb:
        last = LAST.get(ws.title)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    if last is not None and SHEETS[ws.title][0] <= c.row <= last:
                        continue
                    c.value = shift_formula(c.value, ws.title, delta)
    # 2) вставка/видалення рядків і заповнення позицій
    for name, (start, gen) in SHEETS.items():
        resize(wb[name], start, n, gen, items)
    # 3) тексти
    total = sum(it["qty"] * it["price"] for it in items)
    tot = {m: sum(it["qty"] * round(it["price"] * m + 1e-9) for it in items) for m in (1.05, 1.08, 1.1)}
    dd, mm, yyyy = a.date.split(".")
    num = f"КА-{dd}{mm}{yyyy}"
    c = wb["рахунок"]["A1"]; c.value = re.sub(r"КА-\d{8}", num, c.value)
    c = wb["накладна"]["D2"]; c.value = re.sub(r"КА-\d{8}", num, c.value)
    gv = lambda s: f"Загальна вартість (ціна): {fmt_money(s)} грн ({words(s)}) без ПДВ."
    for name, s in (("Дрюк", total), ("Нагорнюк", tot[1.05]), ("Лозко", tot[1.08]), ("ТОВ Колотір", tot[1.1])):
        find(wb[name], "Загальна вартість").value = gv(s)
    for name, s, kop in (("рахунок", total, True), ("накладна", total, True),
                         ("пропозиція 1", tot[1.05], True), ("пропозиція 3", tot[1.1], False)):
        find(wb[name], "Сума прописом").value = f"Сума прописом {words(s, kop)}"
    find(wb["пропозиція 2"], "Всього на суму").value = f"Всього на суму: {words(tot[1.08])}"
    wb.save(a.dst)
    print(f"позицій {n}; рахунок {fmt_money(total)}; +5% {fmt_money(tot[1.05])}; "
          f"+8% {fmt_money(tot[1.08])}; +10% {fmt_money(tot[1.1])}")


if __name__ == "__main__":
    main()
